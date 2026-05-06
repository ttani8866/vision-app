from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "AIレベル別スキル定義"

levels = [
    ("Lv0", "AI未活用",
     "AIツールに関する知識がなく、従来の業務手法を用いる段階。AIの存在は知っているが業務での活用経験はない。データ管理は手作業中心（Excel・紙ベース）。",
     "未活用層"),
    ("Lv1", "AIの基本を学び始める",
     "基本的なAIツールを業務の一部で活用開始。ChatGPTやDeepLなどの汎用AIを利用できる。メール作成補助・アイデア出し・要約・簡単な文書作成・シンプルなリサーチが可能。",
     "入門層"),
    ("Lv2", "日常業務でAIを活用",
     "AIを活用して業務の効率化を図り、その効果を実感し始める段階。メール・挨拶文・議事録を効率的に作成できる。簡単なプロンプトを作成し業務の一部を効率化。文章要約・データ整理・リスト作成にAIを活用できる。",
     "入門層"),
    ("Lv3", "業務の効率化を推進",
     "AIを活用してより多くの業務プロセスを効率化し、具体的な成果につなげる段階。提案書・企画書・SNSコンテンツの作成をAIで効率化。営業活動や社内報告にAIを活用し成果を出し始める。社内勉強会を開催しAI活用の知識を共有できる。プレゼン資料の草案や簡易マニュアルを作成できる。",
     "活用層"),
    ("Lv4", "文章生成・情報収集・データ集計が可能。専門的なAIツールを使い始める",
     "AIを活用した効率的な情報収集・データ分析を行い、業務改善にドライブをかけられる。文章生成AIで企画書・プレゼン資料を作成できる。AIで簡単なデータ集計・分析を行いレポートを作成。適切なプロンプトの組み方を理解し実践できている。",
     "活用層"),
    ("Lv5", "基盤生成AIを使いこなし、専門AIも活用できる（T字型）。簡単な自動化ができる",
     "基盤生成AIを使いこなした上で専門的なAIツールを駆使し、業務効率を大幅に向上させられる。文章生成・画像生成・音声認識AIを使い分けられる（マルチモーダル）。SNSコンテンツをAIで生成できる。ルーチン業務の効率化に成功し作業時間を大幅削減。効果的なプロンプトにより完成度の高い資料・企画書を作れる。GPTsやGemのカスタム設定ができる。GASを使った簡単な自動化を実装できる。コンテクストエンジニアリングの概念を理解し、AIへの指示設計ができる。自分の担当業務の非効率なポイントを特定し、AIを使って改善策を設計・実行できる（個人レベルの業務改善）。改善前後の工数を数値で記録し、効果を自分で検証できる。",
     "推進層"),
    ("Lv6", "複数の生成AIを使い分けて戦略的に活用。全社での生成AI導入をリード可能。自動化を業務に実装できる",
     "複数の生成AIを使い分けて戦略的な活用ができ、クライアントワークでも顕著な成果を出している。AIリサーチを顧客分析・提案に日常的に活用できる。社内AI活用プロジェクトに積極的に関わり貢献できる。業務経験の浅いメンバーにAI活用を指導できる。適切なプロンプトでマーケティングリサーチ（市場・競合分析）と戦略策定ができる。AIを活用して一定レベルの事業戦略の立案ができる。GASやGoogle Workspace Studioを活用し、繰り返し業務の自動化ワークフローを業務に組み込める。バイブコーディングによってプロダクトを開発し、デプロイまで完結できる。チーム・部署単位の業務フローを俯瞰し、複数人が関わる非効率プロセスをAIと自動化で改善できる（チームレベルの業務改善）。改善効果を定量化し、上長・クライアントに説明できる。",
     "推進層"),
    ("Lv7", "プロフェッショナルなT字型AI人材として認定される。自動化・エージェント活用を組織に展開できる",
     "幅広いAI活用スキルを持ち、専門的なAIツールを使いこなせる。基盤AIを活用しながら専門的なAIツールを駆使。高度な事業計画書・戦略資料を自在に作成できる。社内でAI活用の体系的な教育を実施できる。複数の自動化ワークフローを設計・運用し、他者への実装指導ができる。APIを活用したサービス間連携を設計できる。バイブコーディングで社内向けツールを継続的に開発・改善し、チーム全体の生産性向上に貢献できる。AIエージェントの概念を理解し、単体エージェントの設計・運用ができる。部門横断の業務改善プロジェクトをAI・自動化を軸に設計・推進できる（組織レベルの業務改善）。業務改善の成果を社内標準プロセスとして定着させ、他部署への横展開ができる。",
     "専門層"),
    ("Lv8", "全社的なAI活用の定着・変革をリード。エージェント型AIの社内基盤を構築できる",
     "AIを活用し、社内のAI導入・推進の中心的な役割を担う。「AIマスター」と呼べる高いレベルで生成AIに関して広範で深い知識を持ち、実践で通用するアウトプットが可能。AI活用の社内研修を企画・実施できる。AIを活用したプロジェクトを全社レベルで推進し、社内の変革をリードできる。複数の自動化・エージェントを組み合わせた業務フロー全体の再設計ができる。社内のAI情報基盤（ナレッジベース・エージェント連携の仕組み）の構築をリードできる。外部ツール・APIを組み合わせた独自のAIシステムを設計・運用できる。",
     "専門層"),
    ("Lv9", "カスタムAIモデルの開発",
     "AIを活用した新たなビジネスモデルを創出できる。自社専用のカスタムAIモデルを開発し業務最適化を実現。AIの活用方法を社外へ発信（セミナー講師・書籍出版など）している。企業のDXを推進し、AI専門家として社外に情報発信している。",
     "リーダー層"),
    ("Lv10", "業界リーダーとしてのAI活用",
     "AIを活用して業界全体に影響を与える。AIを活用し新たなビジネスを創出し業界標準を確立。AIの導入支援・コンサルティングを行い他社へ技術提供できる。AI活用の第一人者として業界団体に属し業界を牽引している。",
     "リーダー層"),
]

headers = ["レベル", "タイトル", "詳細スキル定義", "区分"]
ws.append(headers)

for row in levels:
    ws.append(list(row))

# Styling
font_name = "Yu Gothic"
header_font = Font(name=font_name, bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", start_color="305496")
body_font = Font(name=font_name, size=10)
level_font = Font(name=font_name, bold=True, size=11)

tier_colors = {
    "未活用層": "F2F2F2",
    "入門層": "DDEBF7",
    "活用層": "FFF2CC",
    "推進層": "FCE4D6",
    "専門層": "E2EFDA",
    "リーダー層": "D9E1F2",
}

thin = Side(border_style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col_idx, _ in enumerate(headers, 1):
    c = ws.cell(row=1, column=col_idx)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border

for r in range(2, len(levels) + 2):
    tier = ws.cell(row=r, column=4).value
    fill = PatternFill("solid", start_color=tier_colors.get(tier, "FFFFFF"))
    for col in range(1, 5):
        cell = ws.cell(row=r, column=col)
        cell.font = level_font if col == 1 else body_font
        cell.alignment = Alignment(
            horizontal="center" if col in (1, 4) else "left",
            vertical="top",
            wrap_text=True,
        )
        cell.fill = fill
        cell.border = border

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 38
ws.column_dimensions["C"].width = 90
ws.column_dimensions["D"].width = 12

ws.row_dimensions[1].height = 28
for r in range(2, len(levels) + 2):
    ws.row_dimensions[r].height = 110

ws.freeze_panes = "A2"

wb.save(r"C:\claude code\AIレベル別スキル定義.xlsx")
print("saved")
